from __future__ import annotations
import io
from dataclasses import dataclass
from pathlib import Path
from opsproof.core.config import MAX_FILE_BYTES, SUPPORTED_EXTENSIONS
from opsproof.core.exceptions import IngestionError

@dataclass(frozen=True)
class IngestedLine:
    text: str
    line_number: int | None = None
    paragraph_number: int | None = None

@dataclass(frozen=True)
class IngestedDocument:
    file_name: str
    kind: str
    lines: list[IngestedLine]

class IngestionService:
    """Decode bounded uploads without persisting or executing their contents."""
    def ingest(self, file_name: str, data: bytes) -> IngestedDocument:
        name = Path(file_name).name
        suffix = Path(name).suffix.lower()
        if suffix not in SUPPORTED_EXTENSIONS:
            raise IngestionError(f"Unsupported file type: {suffix or '(none)'}")
        if not data:
            raise IngestionError("The uploaded runbook is empty.")
        if len(data) > MAX_FILE_BYTES:
            raise IngestionError(f"File exceeds the {MAX_FILE_BYTES} byte limit.")
        if suffix == ".docx":
            return self._docx(name, data)
        if b"\x00" in data:
            raise IngestionError("Binary or unreadable content was rejected.")
        try:
            text = data.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise IngestionError("Text must be valid UTF-8.") from exc
        if not text.strip():
            raise IngestionError("The uploaded runbook contains no readable text.")
        return IngestedDocument(name, suffix[1:], [IngestedLine(x, i) for i, x in enumerate(text.splitlines(), 1)])

    def _docx(self, name: str, data: bytes) -> IngestedDocument:
        # python-docx reads document XML only; macros and embedded objects are never invoked.
        try:
            from docx import Document
            doc = Document(io.BytesIO(data))
        except Exception as exc:
            raise IngestionError("Malformed or unreadable DOCX file.") from exc
        lines = [IngestedLine(p.text, paragraph_number=i) for i, p in enumerate(doc.paragraphs, 1) if p.text.strip()]
        paragraph = len(doc.paragraphs)
        for table_no, table in enumerate(doc.tables, 1):
            for row_no, row in enumerate(table.rows, 1):
                paragraph += 1
                text = " | ".join(cell.text.strip() for cell in row.cells)
                if text.strip(" |"):
                    lines.append(IngestedLine(f"[Table {table_no}, row {row_no}] {text}", paragraph_number=paragraph))
        if not lines:
            raise IngestionError("The DOCX contains no readable paragraphs or table text.")
        return IngestedDocument(name, "docx", lines)
